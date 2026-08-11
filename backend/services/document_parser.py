"""文档解析器：按扩展名分发，提取全文文本。

支持类型：
- 文本类（txt/md/csv/json/html + 代码/配置文件）：标准库 utf-8 读取（errors='replace'，容错乱码）
- pdf：PyMuPDF(fitz) 逐页 get_text，无文本页跳过
- docx：python-docx 段落 + 表格单元格
- xlsx：openpyxl 每 sheet 单元格文本
- pptx：python-pptx 每页 shape 文本

解析失败抛出 ValueError（带友好中文信息），由上层转为 4xx 响应。
"""
import logging

logger = logging.getLogger(__name__)

# 常见代码/配置文件扩展名（映射到 _parse_text 纯文本解析）
_CODE_EXTS = {
    # 代码
    "py", "js", "mjs", "cjs", "jsx", "ts", "tsx", "java", "go", "rs",
    "c", "h", "cpp", "hpp", "cc", "cs", "php", "rb", "swift", "kt",
    "sh", "bash", "zsh", "fish", "ps1", "sql", "lua", "r", "pl",
    "scala", "dart", "vue", "svelte",
    # 配置/标记
    "yaml", "yml", "toml", "ini", "conf", "cfg", "xml", "properties", "env",
}


def _parse_text(path: str) -> str:
    """txt/md/csv/json/html：标准库直接读取，utf-8 容错。"""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    return text.lstrip("\ufeff")  # 去掉可能存在的 UTF-8 BOM


def _parse_pdf(path: str) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise ValueError("服务器缺少 PyMuPDF 库，无法解析 PDF 文件")
    try:
        doc = fitz.open(path)
    except Exception as e:
        raise ValueError(f"PDF 文件解析失败（文件可能损坏或不是有效 PDF）：{e}")
    try:
        if doc.page_count > 500:
            raise ValueError(f"PDF 页数过多（{doc.page_count} 页，上限 500），已拒绝解析")
        pages = []
        for page in doc:
            text = page.get_text()
            if text and text.strip():
                pages.append(text.strip())
        return "\n\n".join(pages)
    finally:
        doc.close()


def _parse_docx(path: str) -> str:
    try:
        from docx import Document
    except ImportError:
        raise ValueError("服务器缺少 python-docx 库，无法解析 docx 文件")
    try:
        doc = Document(path)
    except Exception as e:
        raise ValueError(f"docx 文件解析失败（文件可能损坏或不是有效 Word 文档）：{e}")
    parts = []
    for para in doc.paragraphs:
        if para.text and para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)


def _parse_xlsx(path: str) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError("服务器缺少 openpyxl 库，无法解析 xlsx 文件")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"xlsx 文件解析失败（文件可能损坏或不是有效 Excel 工作簿）：{e}")
    try:
        parts = []
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                parts.append(f"[工作表: {ws.title}]")
                parts.extend(rows)
        return "\n".join(parts)
    finally:
        wb.close()


def _parse_pptx(path: str) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        raise ValueError("服务器缺少 python-pptx 库，无法解析 pptx 文件")
    try:
        prs = Presentation(path)
    except Exception as e:
        raise ValueError(f"pptx 文件解析失败（文件可能损坏或不是有效 PowerPoint 演示文稿）：{e}")
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        slide_texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs).strip()
                    if t:
                        slide_texts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    line = " | ".join(c for c in cells if c)
                    if line:
                        slide_texts.append(line)
        if slide_texts:
            parts.append(f"[第{i}页]")
            parts.extend(slide_texts)
    return "\n".join(parts)


_PARSERS = {
    "txt": _parse_text,
    "md": _parse_text,
    "csv": _parse_text,
    "json": _parse_text,
    "html": _parse_text,
    "pdf": _parse_pdf,
    "docx": _parse_docx,
    "xlsx": _parse_xlsx,
    "pptx": _parse_pptx,
}
# 代码/配置文件：全部走纯文本解析
for _ext in _CODE_EXTS:
    _PARSERS.setdefault(_ext, _parse_text)


def parse_file(path: str, ext: str) -> str:
    """按扩展名解析文件，返回提取的全文。

    path: 文件绝对/相对路径；ext: 小写无点扩展名（如 'pdf'、'docx'）。
    解析失败（不支持类型 / 库缺失 / 文件损坏 / 无文本内容）抛出 ValueError。
    """
    ext = (ext or "").strip().lower().lstrip(".")
    parser = _PARSERS.get(ext)
    if not parser:
        raise ValueError(f"不支持的文件类型: {ext or 'unknown'}")
    try:
        text = parser(path)
    except ValueError:
        raise
    except Exception as e:
        logger.exception("[document_parser] parse failed: %s (%s)", path, ext)
        raise ValueError(f"文件解析失败（{ext} 文件可能损坏或内容异常）：{e}")
    text = (text or "").strip()
    if not text:
        raise ValueError("未能从文件中提取到任何文本内容（文件可能为空或为扫描件）")
    return text

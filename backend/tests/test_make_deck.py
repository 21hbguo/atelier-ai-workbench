"""make_deck 高级演示文稿工具（.pptx + .html 双产物）单元测试（tmp_path 沙箱）。

覆盖：套餐门控 / 登录校验 / 路径越界 / 扩展名错误 / 非法 theme / 空 slides / slides 超限 /
正常生成（7 种版式各一页）：python-pptx 读回验证页数与内容、HTML 读回验证各版式类名与
主题色及文本转义（<script> → &lt;script&gt;）；ctx.files 入队 2 条（先 .pptx 后 .html，
filename/url/size/description 正确）；theme=minimal 时 HTML --primary 为 36454F。
测试模式与 test_make_office.py 一致：无 pytest-asyncio，用 asyncio.run 包装协程。
"""
import asyncio
import os
import urllib.parse

import pytest
from pptx import Presentation

from backend import config
from backend.services.agent.context import AgentContext
from backend.services.agent.tools.make_deck import make_deck


def _run(coro):
    return asyncio.run(coro)


def _ctx(user_id=123, file_write=True):
    features = {"file_write": True} if file_write else {}
    return AgentContext(
        session_id=1,
        user_id=user_id,
        extra={"entitlements": {"features": features}},
    )


@pytest.fixture
def ws(tmp_path, monkeypatch):
    """把用户工作区根目录重定向到 pytest tmp_path（workspace 动态读 config 属性）。"""
    monkeypatch.setattr(config, "USER_WORKSPACES_DIR", tmp_path)
    return tmp_path


def _root(ws, user_id=123):
    return ws / f"user_{user_id}"


def _slide_text(slide) -> str:
    """一页内全部文本框文本（\n 连接），供子串断言（多段落要点是同一文本框）。"""
    return "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


LAYOUTS = ("cover", "section", "title_content", "two_column", "quote", "data_callout", "ending")


def _full_args(**overrides):
    """7 种版式各一页的完整演示文稿参数。"""
    args = {
        "filename": "deck.pptx",
        "theme": "executive",
        "slides": [
            {"layout": "cover", "title": "产品发布会", "subtitle": "2026 年度战略报告"},
            {"layout": "section", "title": "第一部分 市场概览"},
            {"layout": "title_content", "title": "市场规模", "bullets": ["要点一", "要点二"], "notes": "演讲备注"},
            {"layout": "two_column", "title": "竞品对比", "bullets": ["左侧要点"], "right_bullets": ["右侧要点"]},
            {"layout": "quote", "quote": "好的设计是尽可能少的设计", "author": "某设计师"},
            {"layout": "data_callout", "stat": 128, "stat_label": "年度增长率"},
            {"layout": "ending", "title": "谢谢观看"},
        ],
    }
    args.update(overrides)
    return args


# ---------- 公共门控 / 参数校验 ----------

def test_make_deck_requires_entitlement(ws):
    result = _run(make_deck(_full_args(), _ctx(file_write=False)))
    assert result == "当前套餐不支持文件写入。"


def test_make_deck_requires_login(ws):
    result = _run(make_deck(_full_args(), _ctx(user_id=None)))
    assert result == "需要登录后才能使用文件工具"


def test_make_deck_empty_filename(ws):
    result = _run(make_deck(_full_args(filename="   "), _ctx()))
    assert "filename 不能为空" in result


def test_make_deck_outside_rejected(ws):
    result = _run(make_deck(_full_args(filename="../a.pptx"), _ctx()))
    assert "路径无效" in result
    assert "越界" in result


def test_make_deck_wrong_ext(ws):
    result = _run(make_deck(_full_args(filename="a.xls"), _ctx()))
    assert result == "文件名必须以 .pptx 结尾"


def test_make_deck_invalid_theme(ws):
    result = _run(make_deck(_full_args(theme="rainbow"), _ctx()))
    assert "theme 必须是" in result


def test_make_deck_empty_slides(ws):
    result = _run(make_deck(_full_args(slides=[]), _ctx()))
    assert "slides 不能为空" in result


def test_make_deck_slides_over_limit(ws):
    slides = [{"title": f"P{i}"} for i in range(31)]
    result = _run(make_deck(_full_args(slides=slides), _ctx()))
    assert "参数超限" in result
    assert "30" in result


def test_make_deck_bullets_over_limit(ws):
    slides = [{"layout": "title_content", "title": "T", "bullets": [f"b{i}" for i in range(9)]}]
    result = _run(make_deck(_full_args(slides=slides), _ctx()))
    assert "参数超限" in result
    assert "8" in result


# ---------- 正常生成：双产物 + 入队 ctx.files ----------

def test_make_deck_success_enqueues(ws):
    ctx = _ctx()
    result = _run(make_deck(_full_args(), ctx))
    assert "已生成演示文稿【deck.pptx】及网页版预览【deck.html】" in result
    assert "共 7 页" in result
    # 两个文件都已写入工作区
    assert (_root(ws) / "deck.pptx").exists()
    assert (_root(ws) / "deck.html").exists()
    # ctx.files 入队 2 条：先 .pptx 后 .html，url/description 正确
    assert len(ctx.files) == 2
    f0, f1 = ctx.files
    assert f0["filename"] == "deck.pptx"
    assert f0["url"] == f"/api/workspace/files/download?path={urllib.parse.quote('deck.pptx')}"
    assert f0["size"] == os.path.getsize(_root(ws) / "deck.pptx")
    assert f0["description"] == "PPT 演示文稿（可编辑）"
    assert f1["filename"] == "deck.html"
    assert f1["url"] == f"/api/workspace/files/download?path={urllib.parse.quote('deck.html')}"
    assert f1["size"] == os.path.getsize(_root(ws) / "deck.html")
    assert f1["description"] == "HTML 网页版预览"
    # python-pptx 读回验证：7 页、各版式标题/要点/备注
    prs = Presentation(str(_root(ws) / "deck.pptx"))
    assert len(prs.slides) == 7
    slide_texts = [_slide_text(slide) for slide in prs.slides]
    assert "产品发布会" in slide_texts[0]  # cover：标题 + 副标题
    assert "2026 年度战略报告" in slide_texts[0]
    assert "第一部分 市场概览" in slide_texts[1]  # section
    assert "市场规模" in slide_texts[2]  # title_content
    # 要点合并在同一文本框（多段落），用子串匹配
    assert "要点一" in slide_texts[2] and "要点二" in slide_texts[2]
    assert prs.slides[2].notes_slide.notes_text_frame.text == "演讲备注"
    assert "左侧要点" in slide_texts[3] and "右侧要点" in slide_texts[3]  # two_column
    assert "好的设计是尽可能少的设计" in slide_texts[4]  # quote
    assert "某设计师" in slide_texts[4]
    assert "128" in slide_texts[5]  # data_callout（stat 数字转字符串）
    assert "年度增长率" in slide_texts[5]
    assert "谢谢观看" in slide_texts[6]  # ending
    # HTML 读回：各版式类名 + 主题色 + 页码
    html_text = (_root(ws) / "deck.html").read_text(encoding="utf-8")
    for cls in LAYOUTS:
        assert f'class="slide layout-{cls}"' in html_text
    assert "--primary: #1E2761" in html_text  # executive 主题色
    assert "Microsoft YaHei" in html_text
    assert "<div class=\"page-num\">7</div>" in html_text


def test_make_deck_html_escapes_script(ws):
    ctx = _ctx()
    slides = [
        {"layout": "cover", "title": "<script>alert(1)</script>", "subtitle": "<b>加粗</b>"},
        {"layout": "title_content", "title": "T", "bullets": ["<script>alert(2)</script>", "a & b"]},
    ]
    _run(make_deck(_full_args(slides=slides), ctx))
    html_text = (_root(ws) / "deck.html").read_text(encoding="utf-8")
    # 注入脚本被转义为 &lt;script&gt;，原始 <script> 标签不出现
    assert "&lt;script&gt;" in html_text
    assert "&lt;b&gt;加粗&lt;/b&gt;" in html_text
    assert "a &amp; b" in html_text
    assert "<script>alert" not in html_text
    # 原始内容（未转义）仍写入 pptx（可编辑文件保留原文本）
    prs = Presentation(str(_root(ws) / "deck.pptx"))
    assert "<script>alert(1)</script>" in _slide_text(prs.slides[0])


def test_make_deck_theme_minimal_primary(ws):
    ctx = _ctx()
    _run(make_deck(_full_args(theme="minimal"), ctx))
    html_text = (_root(ws) / "deck.html").read_text(encoding="utf-8")
    assert "--primary: #36454F" in html_text


def test_make_deck_subdir_relative_urls(ws):
    """子目录路径：两份文件都在子目录，url 保留路径层级。"""
    ctx = _ctx()
    _run(make_deck(_full_args(filename="docs/report.pptx"), ctx))
    assert (_root(ws) / "docs" / "report.pptx").exists()
    assert (_root(ws) / "docs" / "report.html").exists()
    assert ctx.files[0]["url"] == f"/api/workspace/files/download?path={urllib.parse.quote('docs/report.pptx')}"
    assert ctx.files[1]["url"] == f"/api/workspace/files/download?path={urllib.parse.quote('docs/report.html')}"


# ---------- 注册表 ----------

def test_make_deck_registered_in_registry():
    from backend.services.agent.registry import get_tool

    tool = get_tool("make_deck")
    assert tool is not None
    assert "filename" in tool["parameters"]["required"]
    assert "slides" in tool["parameters"]["required"]
    layouts = tool["parameters"]["properties"]["slides"]["items"]["properties"]["layout"]["enum"]
    assert set(layouts) == set(LAYOUTS)
